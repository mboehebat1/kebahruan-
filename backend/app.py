"""MacroQuest server — Flask + SQLite. Jalankan: pip install -r requirements.txt && python app.py"""
import os, io, json, time, hmac, secrets, sqlite3
from functools import wraps
from flask import Flask, request, jsonify, send_from_directory, send_file, g
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DB = os.environ.get('MQ_DB', os.path.join(HERE, 'macroquest.db'))
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024
SCHEMA = """
CREATE TABLE IF NOT EXISTS users(nim TEXT PRIMARY KEY, name TEXT, kelas TEXT, pass TEXT, role TEXT DEFAULT 'mhs');
CREATE TABLE IF NOT EXISTS progress(nim TEXT PRIMARY KEY, state TEXT, xp INT DEFAULT 0, lv INT DEFAULT 0, score INT DEFAULT 0, bd TEXT DEFAULT '', updated REAL);
CREATE TABLE IF NOT EXISTS sessions(token TEXT PRIMARY KEY, nim TEXT, created REAL);
CREATE TABLE IF NOT EXISTS config(k TEXT PRIMARY KEY, v TEXT);
"""
UPSERT = "INSERT INTO users(nim,name,kelas,pass,role) VALUES(?,?,?,?,'mhs') ON CONFLICT(nim) DO UPDATE SET name=excluded.name,kelas=excluded.kelas,pass=excluded.pass"


def init_db():
    fresh = not os.path.exists(DB)
    d = sqlite3.connect(DB, timeout=10)
    d.executescript('PRAGMA journal_mode=WAL;' + SCHEMA)
    if not d.execute("SELECT 1 FROM users WHERE role='dsn'").fetchone():
        d.execute("INSERT INTO users VALUES('dosen','Dosen','-',?,'dsn')",
                  (generate_password_hash(os.environ.get('MQ_DOSEN_PASS', 'dosen123')),))
    if fresh:  # akun demo hanya saat database baru dibuat
        d.executemany("INSERT INTO users VALUES(?,?,?,?,'mhs')",
                      [('1001', 'Ahmad Contoh', 'A', 'ahmad1'), ('1002', 'Budi Contoh', 'A', 'budi22')])
    d.commit(); d.close()


def db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB, timeout=10); g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop('db', None)
    if d: d.close()


def clean(x): return str(x if x is not None else '').strip()
def body(): return request.get_json(force=True, silent=True) or {}


def auth(role=None):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            t = request.headers.get('Authorization', '').replace('Bearer ', '')
            u = db().execute('SELECT u.* FROM sessions s JOIN users u ON u.nim=s.nim WHERE s.token=? AND s.created>?',
                             (t, time.time() - 30 * 86400)).fetchone()
            if not u: return jsonify(error='Sesi berakhir, silakan login ulang'), 401
            if role and u['role'] != role: return jsonify(error='Khusus dosen'), 403
            g.user = u
            return f(*a, **k)
        return w
    return deco


def get_cfg():
    r = db().execute("SELECT v FROM config WHERE k='config'").fetchone()
    return json.loads(r['v']) if r else {'tasks': {}, 'sched': {}, 'mat': {}, 'bank': None}


def payload(u, token=None):
    p = db().execute('SELECT state FROM progress WHERE nim=?', (u['nim'],)).fetchone()
    return {'token': token, 'now': time.time() * 1000, 'config': get_cfg(),
            'user': {k: u[k] for k in ('nim', 'name', 'kelas', 'role')},
            'state': json.loads(p['state']) if p and p['state'] else None}


@app.post('/api/login')
def login():
    b = body()
    u = db().execute('SELECT * FROM users WHERE nim=?', (clean(b.get('nim')),)).fetchone()
    pw = clean(b.get('pass'))
    ok = u and (check_password_hash(u['pass'], pw) if u['role'] == 'dsn' else hmac.compare_digest(u['pass'].encode(), pw.encode()))
    if not ok: return jsonify(error='NIM atau passcode salah'), 401
    t = secrets.token_hex(24)
    db().execute('INSERT INTO sessions VALUES(?,?,?)', (t, u['nim'], time.time())); db().commit()
    return jsonify(payload(u, t))


@app.get('/api/session')
@auth()
def session(): return jsonify(payload(g.user))


@app.get('/api/state')
@auth()
def get_state():
    return jsonify(state=payload(g.user)['state'])


@app.put('/api/state')
@auth()
def put_state():
    if g.user['role'] == 'dsn': return jsonify(ok=True)
    b = body(); st = b.get('state')
    if not isinstance(st, dict): return jsonify(error='State tidak valid'), 400
    d = db(); nim = g.user['nim']
    row = d.execute('SELECT state FROM progress WHERE nim=?', (nim,)).fetchone()
    old = json.loads(row['state']) if row and row['state'] else {}
    for k, sb in (st.get('sub') or {}).items():  # nilai & feedback dosen tidak boleh tertimpa klien
        o = (old.get('sub') or {}).get(k)
        if o and o.get('v') == sb.get('v'): sb['nilai'], sb['fb'] = o.get('nilai'), o.get('fb', '')
    st['nim'] = nim
    n = lambda k: int(b.get(k) or 0)
    d.execute('INSERT INTO progress(nim,state,xp,lv,score,bd,updated) VALUES(?,?,?,?,?,?,?) ON CONFLICT(nim) DO UPDATE SET '
              'state=excluded.state,xp=excluded.xp,lv=excluded.lv,score=excluded.score,bd=excluded.bd,updated=excluded.updated',
              (nim, json.dumps(st), n('xp'), n('lv'), n('score'), clean(b.get('bd'))[:40], time.time()))
    d.commit(); return jsonify(ok=True)


@app.get('/api/leaderboard')
@auth()
def leaderboard():
    r = db().execute("SELECT u.nim,u.name,COALESCE(p.xp,0) xp,COALESCE(p.lv,0) lv,COALESCE(p.score,0) score,COALESCE(p.bd,'') bd "
                     "FROM users u LEFT JOIN progress p ON p.nim=u.nim WHERE u.role='mhs' ORDER BY xp DESC,score DESC,u.name").fetchall()
    return jsonify([dict(x) for x in r])


@app.put('/api/config')
@auth('dsn')
def put_cfg():
    b = body()
    cfg = {k: b.get(k) for k in ('tasks', 'sched', 'mat', 'bank')}
    for k in ('tasks', 'sched', 'mat'):
        if not isinstance(cfg[k], dict): cfg[k] = {}
    db().execute("INSERT INTO config VALUES('config',?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (json.dumps(cfg),))
    db().commit(); return jsonify(ok=True)


@app.get('/api/roster')
@auth('dsn')
def roster():
    r = db().execute("SELECT u.nim,u.name,u.kelas,u.pass,p.xp,p.lv FROM users u LEFT JOIN progress p ON p.nim=u.nim "
                     "WHERE u.role='mhs' ORDER BY u.nim").fetchall()
    return jsonify([dict(x) for x in r])


@app.post('/api/roster')
@auth('dsn')
def roster_save():
    b = body(); old, nim, name = clean(b.get('old')), clean(b.get('nim')), clean(b.get('name'))
    if not nim or not name or nim == 'dosen': return jsonify(error='NIM (bukan "dosen") dan nama wajib diisi'), 400
    d = db()
    if old and old != nim:
        if d.execute('SELECT 1 FROM users WHERE nim=?', (nim,)).fetchone(): return jsonify(error='NIM sudah dipakai'), 409
        for t in ('users', 'progress', 'sessions'): d.execute(f'UPDATE {t} SET nim=? WHERE nim=?', (nim, old))
    d.execute(UPSERT, (nim, name, clean(b.get('kelas')), clean(b.get('pass')) or secrets.token_hex(3)))
    d.commit(); return jsonify(ok=True)


@app.post('/api/roster/bulk')
@auth('dsn')
def roster_bulk():
    n = 0
    for r in body() if isinstance(body(), list) else []:
        r = [clean(x) for x in r] + ['', '', '', '']
        if not r[0] or not r[1] or r[0] == 'dosen': continue
        db().execute(UPSERT, (r[0], r[1], r[2], r[3] or secrets.token_hex(3))); n += 1
    db().commit(); return jsonify(n=n)


@app.delete('/api/roster/<nim>')
@auth('dsn')
def roster_del(nim):
    for t in ('users', 'progress', 'sessions'): db().execute(f"DELETE FROM {t} WHERE nim=? AND nim!='dosen'", (nim,))
    db().commit(); return jsonify(ok=True)


@app.post('/api/roster/<nim>/reset')
@auth('dsn')
def roster_reset(nim):
    db().execute('DELETE FROM progress WHERE nim=?', (nim,)); db().commit(); return jsonify(ok=True)


@app.post('/api/dosen/pass')
@auth('dsn')
def dosen_pass():
    p = clean(body().get('pass'))
    if len(p) < 6: return jsonify(error='Passcode dosen minimal 6 karakter'), 400
    db().execute("UPDATE users SET pass=? WHERE nim='dosen'", (generate_password_hash(p),)); db().commit()
    return jsonify(ok=True)


@app.get('/api/states')
@auth('dsn')
def states():
    r = db().execute("SELECT p.state FROM progress p JOIN users u ON u.nim=p.nim WHERE u.role='mhs'").fetchall()
    return jsonify([{'state': json.loads(x['state'])} for x in r if x['state']])


@app.post('/api/grade')
@auth('dsn')
def grade():
    b = body(); nim = clean(b.get('nim'))
    r = db().execute('SELECT state FROM progress WHERE nim=?', (nim,)).fetchone()
    if not r: return jsonify(error='Data mahasiswa tidak ditemukan'), 404
    st = json.loads(r['state']); sb = (st.get('sub') or {}).get(str(b.get('id')))
    if not sb: return jsonify(error='Tugas tidak ditemukan'), 404
    v = b.get('nilai'); sb['nilai'] = v if isinstance(v, (int, float)) else None; sb['fb'] = clean(b.get('fb'))
    db().execute('UPDATE progress SET state=? WHERE nim=?', (json.dumps(st), nim)); db().commit()
    return jsonify(ok=True)


@app.post('/api/questions/parse')
@auth('dsn')
def parse_questions():
    f = request.files.get('file')
    if not f: return jsonify(error='File belum dipilih'), 400
    try: ws = load_workbook(f, read_only=True, data_only=True).active
    except Exception: return jsonify(error='File harus berformat .xlsx'), 400
    c = lambda x: clean(x).replace('|', '/').replace("'", '’')
    items, errs = [], []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = list(r or []) + [None] * 8
        if all(x in (None, '') for x in r): continue
        try: lv = int(float(r[0]))
        except (TypeError, ValueError): errs.append(f'Baris {i}: Level harus angka 1-13'); continue
        q, o, a = c(r[1]), [c(x) for x in r[2:6]], c(r[6]).upper()[:1]
        idx = 'ABCD1234'.find(a) % 4 if a and a in 'ABCD1234' else -1
        if not 1 <= lv <= 13 or not q or not all(o) or idx < 0:
            errs.append(f'Baris {i}: level 1-13, pertanyaan, 4 opsi, dan jawaban (A-D) wajib benar'); continue
        items.append({'l': lv, 's': '|'.join([q, *o, str(idx), c(r[7]) or '-'])})
    return jsonify(items=items, errors=errs)


@app.get('/api/template.xlsx')
def template():
    wb = Workbook(); ws = wb.active; ws.title = 'Soal'
    ws.append(['Level', 'Pertanyaan', 'Opsi A', 'Opsi B', 'Opsi C', 'Opsi D', 'Jawaban (A/B/C/D)', 'Pembahasan'])
    ws.append([2, 'Contoh: multiplier keseimbangan 2 sektor adalah...', '1/(1-b)', '1/b', 'b/(1-b)', '1-b', 'A', 'k = 1/(1 − MPC)'])
    for col, w in zip('ABCDEFGH', [8, 60, 24, 24, 24, 24, 18, 50]): ws.column_dimensions[col].width = w
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name='template_soal.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.get('/')
def index(): return send_from_directory(ROOT, 'index.html')
@app.get('/css/<path:p>')
def css(p): return send_from_directory(os.path.join(ROOT, 'css'), p)
@app.get('/js/<path:p>')
def js(p): return send_from_directory(os.path.join(ROOT, 'js'), p)


init_db()
if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))
